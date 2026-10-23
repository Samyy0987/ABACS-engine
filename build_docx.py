import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_document():
    doc = docx.Document()

    # Page setup - Standard Letter / A4 with 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True

        # Header for subsequent pages
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("ABAC ENTERPRISE ENGINE | SYSTEM ARCHITECTURE DOCUMENT")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(128, 128, 128)

        # Footer for subsequent pages
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun1 = fp.add_run("ABAC Enterprise Engine • Architecture Specification | Page ")
        frun1.font.name = "Calibri"
        frun1.font.size = Pt(8.5)
        frun1.font.color.rgb = RGBColor(128, 128, 128)

    # Palette
    NAVY_HEX = "1B365D"
    NAVY_COLOR = RGBColor(27, 54, 93)
    GRAY_COLOR = RGBColor(100, 100, 100)
    DARK_TEXT = RGBColor(40, 40, 40)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    # ================= PAGE 1: TITLE PAGE =================
    # Top header text on page 1
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p_top.add_run("ABAC ENTERPRISE ENGINE | SYSTEM ARCHITECTURE DOCUMENT")
    r.font.name = "Calibri"
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY_COLOR

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(40)

    # University Logo
    logo_path = "sanjivani_logo.png"
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Inches(3.8))
        p_logo.paragraph_format.space_after = Pt(20)

    # University Details
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_uni.paragraph_format.space_after = Pt(4)
    r_uni = p_uni.add_run("SANJIVANI UNIVERSITY, KOPARGAON\n")
    r_uni.font.name = "Calibri"
    r_uni.font.size = Pt(15)
    r_uni.font.bold = True
    r_uni.font.color.rgb = DARK_TEXT

    r_school = p_uni.add_run("(School of Engineering & Technology)\n")
    r_school.font.name = "Calibri"
    r_school.font.size = Pt(13)
    r_school.font.italic = True
    r_school.font.color.rgb = DARK_TEXT

    r_prog = p_uni.add_run("B.Tech – Cyber Security")
    r_prog.font.name = "Calibri"
    r_prog.font.size = Pt(13.5)
    r_prog.font.bold = True
    r_prog.font.color.rgb = DARK_TEXT

    # Spacer
    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(40)

    # Project Title
    p_title_label = doc.add_paragraph()
    p_title_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tl = p_title_label.add_run("Title of the Project")
    r_tl.font.name = "Calibri"
    r_tl.font.size = Pt(13)
    r_tl.font.bold = True
    r_tl.font.color.rgb = DARK_TEXT

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(40)
    r_t = p_title.add_run("“Attribute-Based Access Control (ABAC)\nEngine for Enterprise Applications”")
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(14)
    r_t.font.bold = True
    r_t.font.color.rgb = DARK_TEXT

    # Submitted By
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sb = p_sub.add_run("Submitted By\n")
    r_sb.font.name = "Calibri"
    r_sb.font.size = Pt(12)
    r_sb.font.bold = True
    r_name = p_sub.add_run("Prasad Gaikwad – 2124UCSM1024")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(12)

    # Guide
    p_guide = doc.add_paragraph()
    p_guide.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_guide.paragraph_format.space_before = Pt(25)
    r_g1 = p_guide.add_run("Under the Guidance of\n")
    r_g1.font.name = "Calibri"
    r_g1.font.size = Pt(12)
    r_g1.font.bold = True
    r_g2 = p_guide.add_run("Dr Chetan Bawankar")
    r_g2.font.name = "Calibri"
    r_g2.font.size = Pt(12)
    r_g2.font.bold = True

    # Submitted To
    p_to = doc.add_paragraph()
    p_to.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_to.paragraph_format.space_before = Pt(25)
    r_to1 = p_to.add_run("Submitted To\n")
    r_to1.font.name = "Calibri"
    r_to1.font.size = Pt(11)
    r_to2 = p_to.add_run("Department of Cyber Security\nSanjivani University, Kopargaon")
    r_to2.font.name = "Calibri"
    r_to2.font.size = Pt(12)
    r_to2.font.bold = True

    # Page 1 Footer
    p_foot1 = doc.add_paragraph()
    p_foot1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot1.paragraph_format.space_before = Pt(80)
    r_f1 = p_foot1.add_run("ABAC Enterprise Engine • Architecture Specification | Page 1")
    r_f1.font.name = "Calibri"
    r_f1.font.size = Pt(8.5)
    r_f1.font.color.rgb = GRAY_COLOR

    # ================= PAGE 2: TABLE OF CONTENTS =================
    doc.add_page_break()

    p_toc_head = doc.add_paragraph()
    p_toc_head.paragraph_format.space_before = Pt(20)
    p_toc_head.paragraph_format.space_after = Pt(20)

    # Table of contents table
    toc_data = [
        ("Sr. No.", "Content", "Page No."),
        ("1", "Abstract", "3"),
        ("2", "Introduction", "4"),
        ("3", "Objectives and Scope", "5"),
        ("4", "System Requirements", "6"),
        ("5", "System Architecture", "7"),
        ("6", "Data Flow and Module Design", "9"),
        ("7", "Algorithms and Policy Evaluation Techniques", "10"),
        ("8", "Technology Stack", "11"),
        ("9", "Security and Access Control Considerations", "12"),
        ("10", "Testing and Validation", "13"),
        ("11", "Limitations and Future Scope", "15"),
        ("12", "Conclusion", "16"),
        ("13", "References", "17")
    ]

    table = doc.add_table(rows=len(toc_data), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.0), Inches(4.3), Inches(1.2)]

    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            cell.text = toc_data[row_idx][col_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx != 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            
            run = p.runs[0]
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            if row_idx == 0:
                set_cell_background(cell, NAVY_HEX)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.color.rgb = DARK_TEXT
                if row_idx % 2 == 1:
                    set_cell_background(cell, "F8F9FA")

    # ================= PAGE 3: ABSTRACT =================
    doc.add_page_break()

    h_abs = doc.add_heading(level=1)
    h_abs.paragraph_format.space_before = Pt(10)
    h_abs.paragraph_format.space_after = Pt(12)
    r = h_abs.add_run("ABSTRACT")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = NAVY_COLOR

    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(10)
    p1.paragraph_format.line_spacing = 1.15
    r = p1.add_run(
        "The Attribute-Based Access Control (ABAC) Engine for Enterprise Applications (EasyAccess) is a high-performance, "
        "production-grade authorization framework developed using Python Flask, SQLAlchemy, Oracle REST Data Services (ORDS), "
        "and a modern React TypeScript presentation interface. Its primary purpose is to provide fine-grained, dynamic access "
        "decision capabilities for modern enterprise microservices, zero-trust cloud architectures, and sensitive multi-tenant data platforms."
    )
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(10)
    p2.paragraph_format.line_spacing = 1.15
    r = p2.add_run(
        "Traditional Role-Based Access Control (RBAC) suffers from severe role explosion and lacks contextual awareness. "
        "The EasyAccess ABAC Engine overcomes these limitations by evaluating multi-dimensional authorization policies across four "
        "distinct attribute vectors: Subject attributes (department, clearance tier, assigned roles, MFA verification status), "
        "Resource attributes (classification tiers from Public to Top Secret, resource type, ownership), Action types (read, write, update, "
        "delete, approve, execute), and Environmental conditions (request timestamp, time windows, client IP subnet, device compliance). "
        "Access requests are evaluated against security policies with deterministic Deny-Overrides conflict resolution and strict "
        "Default-Deny enforcement whenever policies are missing or unresolvable."
    )
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)

    p3 = doc.add_paragraph()
    p3.paragraph_format.space_after = Pt(10)
    p3.paragraph_format.line_spacing = 1.15
    r = p3.add_run(
        "The architecture features a resilient dual-tier audit logging system that records every authorization decision with sub-millisecond "
        "telemetry into Oracle ORDS REST endpoints, automatically engaging a local SQLite fallback store during network degradation or database "
        "unavailability to guarantee zero audit trail loss. Sensitive credentials, passwords, and tokens are rigorously sanitized and masked "
        "prior to log persistence. An enterprise administrative web dashboard allows security operators to inspect policy rules, visualize active metrics, "
        "manage policy toggles, and simulate live access requests in real-time."
    )
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)

    p4 = doc.add_paragraph()
    p4.paragraph_format.space_after = Pt(16)
    p4.paragraph_format.line_spacing = 1.15
    r = p4.add_run(
        "The system has been validated through automated testing suites covering policy evaluation logic, credential redaction, database pooling, "
        "and failover persistence. It provides enterprise security teams with a verifiable, NIST SP 800-162 compliant access control engine."
    )
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)

    p_kw = doc.add_paragraph()
    r_kwh = p_kw.add_run("Keywords: ")
    r_kwh.font.name = "Calibri"
    r_kwh.font.size = Pt(10.5)
    r_kwh.font.bold = True
    r_kwt = p_kw.add_run("Attribute-Based Access Control (ABAC), Zero-Trust Architecture, Policy Decision Point (PDP), Policy Enforcement Point (PEP), Deny-Overrides, Oracle ORDS, Python Flask, React, Data Classification, Audit Logging.")
    r_kwt.font.name = "Calibri"
    r_kwt.font.size = Pt(10.5)

    # Save document
    output_filename = "ABAC_Enterprise_Engine_Architecture_Document.docx"
    doc.save(output_filename)
    print(f"Initial docx generated: {output_filename}")

create_document()
